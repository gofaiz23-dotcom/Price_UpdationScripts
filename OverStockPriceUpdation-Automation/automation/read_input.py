"""Read the input workbook: column SKU and column Proposed Price."""

from __future__ import annotations

from pathlib import Path
from typing import Dict

from openpyxl import load_workbook


def load_price_sheet(input_dir: str | Path) -> Dict[str, float]:
    folder = Path(input_dir)
    files = sorted(folder.glob("*.xlsx"))
    files = [p for p in files if not p.name.startswith("~$")]
    if not files:
        raise FileNotFoundError(f"Put the price sheet (.xlsx) in {folder}")

    path = files[0]
    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError(f"{path.name} is empty")

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    try:
        sku_idx = header.index("sku")
        price_idx = header.index("proposed price")
    except ValueError as exc:
        raise ValueError(
            f"{path.name} must have columns SKU and Proposed Price. Found: {rows[0]}"
        ) from exc

    prices: Dict[str, float] = {}
    for row_number, row in enumerate(rows[1:], start=2):
        sku_raw = row[sku_idx] if sku_idx < len(row) else None
        price_raw = row[price_idx] if price_idx < len(row) else None
        sku = "" if sku_raw is None else str(sku_raw).strip()
        if not sku and (price_raw is None or str(price_raw).strip() == ""):
            continue
        if not sku:
            raise ValueError(f"Row {row_number} is missing SKU")
        if price_raw is None or str(price_raw).strip() == "":
            raise ValueError(f"Row {row_number} SKU {sku} is missing Proposed Price")
        try:
            price = float(price_raw)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"Row {row_number} SKU {sku} has a non-numeric price") from exc
        if price < 0:
            raise ValueError(f"Row {row_number} SKU {sku} has a negative price")
        if sku in prices:
            raise ValueError(f"Duplicate SKU in {path.name}: {sku}")
        prices[sku] = price

    if not prices:
        raise ValueError(f"{path.name} has no price rows")
    return prices
