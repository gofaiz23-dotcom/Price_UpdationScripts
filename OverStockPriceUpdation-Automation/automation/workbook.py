"""Parse Overstock Cost Export Manage Cost sheet (openpyxl).

Overstock exports often ship invalid stylesheet colors (not aRGB hex). openpyxl
rejects those; we sanitize / strip styles before reading cell values.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

logger = logging.getLogger(__name__)

_RGB_ATTR_RE = re.compile(r'\brgb="([^"]*)"', re.I)
_HEX_RE = re.compile(r"^[A-Fa-f0-9]{6}$|^[A-Fa-f0-9]{8}$")

# Minimal styles.xml so openpyxl can open workbooks after we drop broken styles.
_MINIMAL_STYLES_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <fonts count="1"><font><sz val="11"/><color theme="1"/><name val="Calibri"/><family val="2"/></font></fonts>
  <fills count="1"><fill><patternFill patternType="none"/></fill></fills>
  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/></cellXfs>
  <cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>
</styleSheet>
"""


def _fix_rgb_attr(match: re.Match[str]) -> str:
    raw = (match.group(1) or "").strip().lstrip("#")
    if _HEX_RE.fullmatch(raw):
        if len(raw) == 6:
            raw = "00" + raw
        return f'rgb="{raw}"'
    return 'rgb="00000000"'


def _sanitize_styles_xml(data: bytes) -> bytes:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("utf-8", errors="ignore")
    fixed = _RGB_ATTR_RE.sub(_fix_rgb_attr, text)
    return fixed.encode("utf-8")


def _xlsx_bytes_with_fixed_styles(file_path: Path, *, replace_styles: bool = False) -> bytes:
    """Return xlsx bytes with styles.xml sanitized or replaced."""
    out = io.BytesIO()
    with zipfile.ZipFile(file_path, "r") as zin, zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            payload = zin.read(info.filename)
            if info.filename == "xl/styles.xml":
                if replace_styles:
                    payload = _MINIMAL_STYLES_XML.encode("utf-8")
                else:
                    payload = _sanitize_styles_xml(payload)
            zout.writestr(info, payload)
    return out.getvalue()


def _is_stylesheet_error(exc: BaseException) -> bool:
    msg = str(exc).lower()
    return "stylesheet" in msg or "argb" in msg or "color" in msg or "invalid xml" in msg


def load_workbook_tolerant(
    file_path: str | Path,
    *,
    read_only: bool = True,
    data_only: bool = True,
):
    """Load Overstock Cost Export; tolerate broken stylesheet colors / XML."""
    path = Path(file_path)
    try:
        return load_workbook(filename=str(path), read_only=read_only, data_only=data_only)
    except (ValueError, TypeError) as exc:
        if not _is_stylesheet_error(exc):
            raise
        logger.warning(
            "openpyxl rejected Overstock stylesheet (%s) — retrying with sanitized colors",
            exc,
        )

    try:
        sanitized = _xlsx_bytes_with_fixed_styles(path, replace_styles=False)
        return load_workbook(
            filename=io.BytesIO(sanitized), read_only=read_only, data_only=data_only
        )
    except (ValueError, TypeError) as exc:
        if not _is_stylesheet_error(exc):
            raise
        logger.warning(
            "Sanitized styles still invalid (%s) — retrying with minimal stylesheet",
            exc,
        )

    minimal = _xlsx_bytes_with_fixed_styles(path, replace_styles=True)
    return load_workbook(
        filename=io.BytesIO(minimal), read_only=read_only, data_only=data_only
    )


# Back-compat alias used by older call sites / tests.
_load_workbook_tolerant = load_workbook_tolerant
